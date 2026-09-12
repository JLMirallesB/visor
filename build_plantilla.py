#!/usr/bin/env python3
"""Genera plantilla-hojas.xlsx: las hojas Ficha, Informes, Anuladas y Listas
listas para trasplantar al libro que crea Microsoft Forms.

Las cabeceras se importan de build_demo.py para que no puedan desincronizarse
de las del archivo de ejemplo.

Las listas cerradas viven en la hoja Listas y las validaciones apuntan a sus
rangos, en lugar de llevar los valores escritos dentro de la regla. Así cada
centro recorta o amplía la lista insertando y borrando filas enteras, sin tocar
ninguna validación.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from build_demo import FICHA_HEADERS, INFORME_HEADERS

OUT = Path(__file__).parent / "plantilla-hojas.xlsx"

# Hoja Anuladas: Id de las actuaciones que el visor deja de mostrar. Se anula por
# Id y no borrando la fila de Sheet1, porque el vínculo con Forms solo añade y
# restaura las filas borradas con la siguiente respuesta.
ANULADAS_HEADERS = ["Id", "Motivo", "Fecha"]

# Hasta qué fila llega la validación en la hoja Ficha. El desplegable es una
# propiedad de las celdas, no de la cabecera.
ULTIMA_FILA = 1000

# Especialidades de las enseñanzas profesionales de música.
# Fuente: Decreto 158/2007, art. 6, texto consolidado (incluye las añadidas por
# el D 109/2011, el D 90/2015 y la O 64/2013). Se respeta la grafía del decreto,
# que mezcla castellano y valenciano: «Cant valencià», «Violoncello».
ESPECIALIDADES_MUSICA = [
    "Acordeón", "Arpa", "Bajo eléctrico", "Cante flamenco", "Canto",
    "Cant valencià", "Clarinete", "Clave", "Contrabajo", "Dulzaina", "Fagot",
    "Flabiol y tamboril", "Flauta travesera", "Flauta de pico", "Gaita",
    "Guitarra", "Guitarra eléctrica", "Guitarra flamenca",
    "Instrumentos de cuerda pulsada del Renacimiento y del Barroco",
    "Instrumentos de púa", "Oboe", "Órgano", "Percusión", "Piano", "Saxofón",
    "Tenora", "Tible", "Trombón", "Trompa", "Trompeta", "Tuba", "Txistu",
    "Viola", "Viola de Gamba", "Violín", "Violoncello",
]

# Especialidades de las enseñanzas profesionales de danza.
# Fuente: Decreto 156/2007, art. 6.
ESPECIALIDADES_DANZA = [
    "Baile flamenco", "Danza clásica", "Danza contemporánea", "Danza española",
]

# «Danza» a secas, que no está en ningún decreto, para el alumnado de las
# enseñanzas ELEMENTALES de danza: el Decreto 157/2007 no fija especialidades
# para ese nivel, solo asignaturas (art. 5), así que no hay ninguna que elegir.
# En música no hace falta la equivalente: las elementales sí tienen especialidad
# desde primero, que es el instrumento.
ESPECIALIDAD_GENERICA = ["Danza"]

# Cada lista cerrada: columna de la hoja Listas → (rótulo, valores, columnas de
# Ficha que la usan).
LISTAS = [
    ("Curso", ["1 EEM", "2 EEM", "3 EEM", "4 EEM", "1 EPM",
               "2 EPM", "3 EPM", "4 EPM", "5 EPM", "6 EPM"],
     ["Curso"]),
    ("Especialidad", ESPECIALIDADES_MUSICA + ESPECIALIDAD_GENERICA + ESPECIALIDADES_DANZA,
     ["Especialidad"]),
    ("Seguimiento", ["Prioritario", "Activo", "Pausa", "Cerrado"],
     ["Seguimiento"]),
    ("Situación", ["Espera de actuación/respuesta nuestra",
                   "En espera de actuación/respuesta de otros",
                   "Ninguno"],
     ["Situación"]),
    ("Estado UEO / ERG", ["Sí", "No", "En Proceso"],
     ["Activada UEO?", "Autorización para Contactar ERG?", "Contactado ERG?"]),
]


def add_headers(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        ancho = max(12, min(40, len(str(cell.value)) + 2))
        ws.column_dimensions[cell.column_letter].width = ancho
    ws.freeze_panes = "A2"


def build_listas(wb):
    ws = wb.create_sheet("Listas")
    add_headers(ws, [rotulo for rotulo, _, _ in LISTAS])
    rangos = {}
    for i, (rotulo, valores, columnas_ficha) in enumerate(LISTAS, start=1):
        letra = get_column_letter(i)
        for fila, valor in enumerate(valores, start=2):
            ws.cell(row=fila, column=i, value=valor)
        ref = f"Listas!${letra}$2:${letra}${len(valores) + 1}"
        for col in columnas_ficha:
            rangos[col] = ref
    return ws, rangos


def add_validaciones(ws, headers, rangos):
    for columna, ref in rangos.items():
        letra = get_column_letter(headers.index(columna) + 1)
        dv = DataValidation(
            type="list",
            formula1=ref,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Valor no admitido",
            error="Elige uno de los valores de la lista (hoja «Listas»).",
        )
        # No se toca showDropDown: en openpyxl ese atributo está invertido y
        # ponerlo en True es justo lo que esconde la flecha del desplegable.
        ws.add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{ULTIMA_FILA}")


def main():
    wb = Workbook()
    ws_fic = wb.active
    ws_fic.title = "Ficha"
    add_headers(ws_fic, FICHA_HEADERS)
    add_headers(wb.create_sheet("Informes"), INFORME_HEADERS)
    add_headers(wb.create_sheet("Anuladas"), ANULADAS_HEADERS)
    _, rangos = build_listas(wb)
    add_validaciones(ws_fic, FICHA_HEADERS, rangos)
    wb.save(OUT)
    print(f"Generado: {OUT}")
    print(f"  Ficha:    {len(FICHA_HEADERS)} columnas, "
          f"{len(rangos)} con desplegable hasta la fila {ULTIMA_FILA}")
    print(f"  Informes: {len(INFORME_HEADERS)} columnas")
    print(f"  Anuladas: {len(ANULADAS_HEADERS)} columnas")
    print(f"  Listas:   {len(LISTAS)} listas — "
          f"{len(ESPECIALIDADES_MUSICA)} especialidades de música + "
          f"{len(ESPECIALIDADES_DANZA)} de danza + «Danza» genérica")


if __name__ == "__main__":
    main()
