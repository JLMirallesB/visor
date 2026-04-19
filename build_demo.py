#!/usr/bin/env python3
"""Genera un Excel de demostración para el visor."""
from datetime import datetime, timedelta
from openpyxl import Workbook

OUT = "/Users/miralles/Documents/GitHub/visor/demo.xlsx"
TODAY = datetime(2026, 4, 11)

# ---------------------------------------------------------------------------
# Hoja "Ficha"
# ---------------------------------------------------------------------------
FICHA_HEADERS = [
    "NIA","Nombre","Apellidos","Fecha de nacimiento","Curso","Especialidad",
    "Centro ERG","Teléfono Alumno","Mail Alumno",
    "Nombre Madre","Teléfono Madre","Mail Madre",
    "Nombre Padre","Teléfono Padre","Mail Padre",
    "Tutor/a","Diagnóstico","Seguimiento","Situación","Campo Abierto",
    "Activada UEO?","Autorización para Contactar ERG?","Contactado ERG?"
]

ficha_rows = [
    ["10001","Lucía","Pérez García",        datetime(2013, 6, 14),"2 EEM","PIANO",
     "IES Joanot Martorell","612345001","lucia.perez@ej.es",
     "Marta García","612345101","marta.garcia@ej.es",
     "Antonio Pérez","612345201","antonio.perez@ej.es",
     "Ana Ferrer","TDAH","Prioritario","Espera de actuación/respuesta nuestra",
     "Necesita refuerzo en autocontrol durante las clases colectivas.",
     "Sí","Sí","Sí"],

    ["10002","Hugo","Martínez López",       datetime(2011, 11, 3),"4 EEM","VIOLÍN",
     "IES Pere Boïl","612345002","hugo.martinez@ej.es",
     "Carmen López","612345102","carmen.lopez@ej.es",
     "Luis Martínez","","",
     "Carlos Mestre","Dislexia","Activo","En espera de actuación/respuesta de otros",
     "Pendiente informe del logopeda externo.",
     "Sí","En Proceso","No"],

    ["10003","Marta","Sánchez Romero",      datetime(2010, 1, 22),"2 EPM","GUITARRA",
     "IES Lluís Vives","612345003","marta.sanchez@ej.es",
     "Pilar Romero","612345103","pilar.romero@ej.es",
     "","","",
     "Elena Torres","","Pausa","Ninguno","",
     "No","No","No"],

    ["10004","Adrián","Fernández Ortiz",    datetime(2012, 9, 8),"3 EEM","FLAUTA",
     "Col·legi La Pureza","612345004","",
     "Rosa Ortiz","612345104","rosa.ortiz@ej.es",
     "Miguel Fernández","612345204","miguel.fernandez@ej.es",
     "Joan Ribó","","Cerrado","Ninguno",
     "Caso cerrado tras adaptación satisfactoria al grupo.",
     "No","No","No"],

    ["10005","Carla","Gómez Ruiz",          datetime(2014, 3, 30),"1 EEM","VIOLONCHELO",
     "IES Joanot Martorell","612345005","carla.gomez@ej.es",
     "Sara Ruiz","612345105","sara.ruiz@ej.es",
     "David Gómez","612345205","david.gomez@ej.es",
     "Ana Ferrer","TDAH, ansiedad","Prioritario","Espera de actuación/respuesta nuestra",
     "Episodios de bloqueo en exámenes. Coordinación con tutora externa.",
     "Sí","Sí","En Proceso"],

    ["10006","Pablo","Díaz Moreno",         datetime(2009, 5, 17),"4 EPM","SAXOFÓN",
     "IES Lluís Vives","612345006","pablo.diaz@ej.es",
     "","","",
     "","","",
     "María Vidal","Altas capacidades","Activo","Ninguno",
     "Programa de enriquecimiento curricular en colaboración con su instituto.",
     "No","No","No"],

    ["10007","Elena","Jiménez Castro",      datetime(2010, 10, 11),"1 EPM","CANTO",
     "Col·legi La Pureza","612345007","elena.jimenez@ej.es",
     "Lucía Castro","612345107","lucia.castro@ej.es",
     "Pedro Jiménez","612345207","",
     "Carlos Mestre","Ansiedad escénica","Activo","En espera de actuación/respuesta de otros",
     "Pendiente cita con psicóloga externa.",
     "En Proceso","No","No"],

    ["10008","Mateo","Álvarez Vega",        datetime(2013, 2, 27),"2 EEM","TROMPETA",
     "IES Pere Boïl","612345008","mateo.alvarez@ej.es",
     "Inés Vega","612345108","ines.vega@ej.es",
     "","","",
     "Joan Ribó","","Pausa","Ninguno","",
     "No","No","No"],

    ["10009","Sofía","Ruiz Navarro",        datetime(2011, 8, 4),"4 EEM","ARPA",
     "IES Joanot Martorell","612345009","sofia.ruiz@ej.es",
     "Beatriz Navarro","612345109","beatriz.navarro@ej.es",
     "José Ruiz","612345209","jose.ruiz@ej.es",
     "Elena Torres","Dislexia, ansiedad","Prioritario","Espera de actuación/respuesta nuestra",
     "Adaptación de partituras y apoyo emocional continuado.",
     "Sí","Sí","Sí"],

    ["10010","Daniel","Moreno Soler",       datetime(2007, 12, 19),"5 EPM","PIANO",
     "IES Lluís Vives","612345010","daniel.moreno@ej.es",
     "Eva Soler","612345110","eva.soler@ej.es",
     "Javier Moreno","612345210","javier.moreno@ej.es",
     "María Vidal","","Activo","Ninguno",
     "Preparando pruebas de acceso al Superior.",
     "No","No","No"],

    ["10011","Inés","Torres Gil",           datetime(2012, 7, 6),"3 EEM","CLARINETE",
     "Col·legi La Pureza","612345011","ines.torres@ej.es",
     "Patricia Gil","612345111","patricia.gil@ej.es",
     "","","",
     "Ana Ferrer","TDAH","Activo","En espera de actuación/respuesta de otros",
     "",
     "Sí","En Proceso","No"],

    ["10012","Marcos","Serrano Bravo",      datetime(2010, 4, 2),"2 EPM","PERCUSIÓN",
     "IES Pere Boïl","612345012","marcos.serrano@ej.es",
     "Lourdes Bravo","612345112","",
     "Andrés Serrano","612345212","andres.serrano@ej.es",
     "Carlos Mestre","","Pausa","Ninguno",
     "En seguimiento por su tutor del centro de referencia.",
     "No","No","No"],

    ["10013","Júlia","Ortega Hidalgo",      datetime(2014, 1, 15),"1 EEM","VIOLA",
     "IES Joanot Martorell","612345013","julia.ortega@ej.es",
     "Cristina Hidalgo","612345113","cristina.hidalgo@ej.es",
     "Raúl Ortega","612345213","raul.ortega@ej.es",
     "Joan Ribó","","Activo","Espera de actuación/respuesta nuestra",
     "Primera matriculación; integración favorable.",
     "No","No","No"],

    ["10014","Aitor","Romero Cano",         datetime(2006, 11, 28),"6 EPM","GUITARRA",
     "","","aitor.romero@ej.es",
     "","","",
     "","","",
     "María Vidal","","Cerrado","Ninguno",
     "Caso cerrado al finalizar enseñanzas profesionales.",
     "No","No","No"],

    ["10015","Noa","Vázquez Peña",          datetime(2013, 9, 21),"2 EEM","FLAUTA DE PICO",
     "Col·legi La Pureza","612345015","noa.vazquez@ej.es",
     "Sonia Peña","612345115","sonia.pena@ej.es",
     "","","",
     "Elena Torres","TEA leve","Prioritario","Espera de actuación/respuesta nuestra",
     "Adaptaciones metodológicas en la clase colectiva. Coordinación con orientadora del IES.",
     "Sí","Sí","En Proceso"],
]

# ---------------------------------------------------------------------------
# Hoja "Sheet1" (Registro)
# ---------------------------------------------------------------------------
REG_HEADERS = [
    "Id","Hora de inicio","Hora de finalización","Correo electrónico","Nombre",
    "NIA","Curso Académico","Tipo de Entrada","Título/Asunto","Descripción",
    "Profesional que registra","Enlace a documento (subir pdf)","Otros"
]

PROFS = [
    ("carmen.albert@ej.es", "Carmen Albert"),
    ("jordi.bosch@ej.es",   "Jordi Bosch"),
    ("sandra.climent@ej.es","Sandra Climent"),
    ("pere.ferrer@ej.es",   "Pere Ferrer"),
]

SAMPLE_PDF = "https://gvaedu-my.sharepoint.com/personal/orientacion/Documents/Apps/Microsoft%20Forms/Registro/Adjuntos/ejemplo.pdf"

def days_ago(n):
    return TODAY - timedelta(days=n)

# (NIA, días_atrás, tipo, título, descripción, prof_idx, con_pdf, otros)
acts = [
    # Lucía 10001 - prioritaria con actividad reciente
    (10001,  3, "Reunión", "Reunión con tutora del IES",
     "Coordinación de pautas de apoyo en aula. Compromiso de revisión en 15 días.", 0, True, ""),
    (10001, 22, "Comunicación familia", "Llamada a familia",
     "Se informa a la madre del avance en autocontrol y se acuerda nueva cita.", 0, False, "Familia receptiva"),
    (10001, 60, "Decisión", "Adaptación metodológica",
     "Se acuerda permitir descansos cortos durante las clases largas.", 1, False, ""),
    (10001,140, "Informe", "Informe de evaluación inicial",
     "Documento entregado al equipo docente con perfil y necesidades.", 1, True, ""),

    # Hugo 10002 - activo
    (10002,  9, "Reunión", "Reunión con logopeda externo",
     "Puesta en común de estrategias de lectoescritura aplicadas en lenguaje musical.", 2, False, ""),
    (10002, 75, "Comunicación familia", "Reunión con familia",
     "Información sobre evolución del trimestre y materiales adaptados.", 2, False, ""),
    (10002,180, "Informe", "Informe inicial de dislexia",
     "Recopilación de antecedentes y recomendaciones del logopeda.", 0, True, ""),

    # Marta 10003 - pausa, sin actividad reciente
    (10003, 95, "Decisión", "Pausar seguimiento",
     "Se decide pausar el seguimiento tras evolución estable.", 3, False, "A revisar en 6 meses"),
    (10003,190, "Reunión", "Reunión con tutora externa",
     "Intercambio de información con su tutora del centro de referencia.", 3, False, ""),

    # Adrián 10004 - cerrado
    (10004,210, "Decisión", "Cierre de caso",
     "Adaptación satisfactoria al grupo. Se cierra el seguimiento activo.", 1, False, ""),

    # Carla 10005 - prioritaria SIN actividad reciente (alerta)
    (10005, 45, "Suceso", "Episodio de bloqueo en examen",
     "Crisis de ansiedad durante la prueba trimestral. Se acompaña y se contacta con familia.", 0, True, ""),
    (10005, 70, "Comunicación familia", "Reunión con madre",
     "Acuerdos sobre rutinas en casa antes de las pruebas.", 0, False, ""),
    (10005,110, "Derivación", "Derivación a psicóloga externa",
     "Se sugiere a la familia consulta con profesional especializado.", 2, True, ""),
    (10005,160, "Informe", "Informe de seguimiento",
     "Documento resumen del primer trimestre.", 0, True, ""),
    (10005,200, "Reunión", "Reunión inicial con tutora",
     "Presentación del caso y necesidades del alumno.", 1, False, ""),

    # Pablo 10006 - activo, altas capacidades
    (10006, 18, "Decisión", "Programa de enriquecimiento",
     "Se acuerda repertorio adicional y proyecto personal.", 3, False, ""),
    (10006, 95, "Reunión", "Coordinación con instituto",
     "Reunión con orientadora del IES para complementar plan.", 3, False, ""),

    # Elena 10007
    (10007, 12, "Comunicación familia", "Llamada a madre",
     "Información sobre próxima cita con psicóloga externa.", 1, False, ""),
    (10007, 50, "Suceso", "Bloqueo en audición",
     "Episodio de ansiedad durante audición. Se acompaña.", 1, True, ""),
    (10007,130, "Informe", "Informe de evaluación",
     "Resumen del primer trimestre.", 2, True, ""),

    # Mateo 10008 - pausa
    (10008,170, "Decisión", "Pausar seguimiento",
     "Sin necesidades urgentes. Se mantiene observación.", 3, False, ""),

    # Sofía 10009 - prioritaria SIN actividad reciente (alerta)
    (10009, 55, "Reunión", "Reunión con familia",
     "Repaso de adaptaciones. Familia colaboradora.", 0, False, ""),
    (10009,100, "Decisión", "Adaptación de partituras",
     "Se proporcionan materiales con tipografía y espaciado adaptados.", 0, True, ""),
    (10009,150, "Informe", "Informe de necesidades",
     "Recopilación de adaptaciones aplicadas.", 0, True, ""),

    # Daniel 10010 - activo, prepara pruebas
    (10010,  6, "Reunión", "Tutoría de orientación académica",
     "Información sobre Conservatorio Superior y plazos.", 2, False, ""),
    (10010, 35, "Comunicación familia", "Reunión informativa",
     "Resolución de dudas sobre pruebas de acceso.", 2, False, ""),
    (10010, 85, "Decisión", "Plan de preparación",
     "Calendario de simulacros y materiales de apoyo.", 3, True, ""),
    (10010,165, "Reunión", "Inicio del seguimiento",
     "Primera entrevista del curso.", 3, False, ""),

    # Inés 10011 - activa
    (10011, 28, "Comunicación familia", "Reunión con madre",
     "Coordinación de pautas en casa.", 0, False, ""),
    (10011,120, "Reunión", "Coordinación con tutora externa",
     "Intercambio de estrategias con tutora del centro de referencia.", 1, False, ""),

    # Marcos 10012 - pausa
    (10012,180, "Decisión", "Pausar seguimiento",
     "Caso pausado tras estabilización.", 1, False, ""),

    # Júlia 10013 - activa, recién matriculada
    (10013, 25, "Reunión", "Acogida inicial",
     "Primera entrevista con la familia y la alumna.", 3, False, ""),
    (10013, 80, "Comunicación familia", "Llamada de seguimiento",
     "Familia satisfecha con la integración.", 3, False, ""),
    (10013,140, "Otro", "Visita guiada al centro",
     "Acompañamiento durante los primeros días.", 2, False, ""),

    # Aitor 10014 - sin actuaciones (caso cerrado)

    # Noa 10015 - prioritaria, actividad reciente
    (10015,  5, "Reunión", "Coordinación con orientadora del IES",
     "Reunión telemática para alinear adaptaciones.", 2, True, ""),
    (10015, 40, "Decisión", "Adaptaciones metodológicas",
     "Material visual estructurado para clases colectivas.", 2, True, ""),
    (10015, 95, "Informe", "Informe inicial",
     "Recopilación de necesidades y antecedentes.", 0, True, ""),
]

def build_registro_rows():
    rows = []
    for i, (nia, d, tipo, titulo, desc, prof_idx, with_pdf, otros) in enumerate(acts, start=1):
        date = days_ago(d)
        end = date + timedelta(minutes=20)
        mail, nombre = PROFS[prof_idx]
        rows.append([
            i,                                      # Id
            date,                                   # Hora de inicio
            end,                                    # Hora de finalización
            mail,                                   # Correo electrónico
            nombre,                                 # Nombre
            nia,                                    # NIA
            "",                                     # Curso Académico (lo calcula la app)
            tipo,                                   # Tipo de Entrada
            titulo,                                 # Título/Asunto
            desc,                                   # Descripción
            "",                                     # Profesional que registra (la app usa Nombre)
            SAMPLE_PDF if with_pdf else "",         # Enlace a documento
            otros,                                  # Otros
        ])
    return rows

# ---------------------------------------------------------------------------
# Hoja "Informes"
# ---------------------------------------------------------------------------
INFORME_HEADERS = [
    "NIA","Fecha Informe","Barreras Acceso","Fortalezas Acceso",
    "Barreras Participación","Fortalezas Participación",
    "Barreras Aprendizaje","Fortalezas Aprendizaje",
    "Necesidades específicas de apoyo educativo","Justificación y orientaciones"
]

informe_rows = [
    # Lucía 10001 — TDAH
    [10001, datetime(2025, 10, 15),
     "Dificultad para mantener la atención durante sesiones largas.",
     "Alta motivación hacia la música y buena relación con sus iguales.",
     "Tiende a aislarse en actividades grupales cuando se siente insegura.",
     "Participación activa cuando se le asigna un rol claro.",
     "Bajo rendimiento en lectura de partituras a primera vista.",
     "Excelente oído y capacidad memorística para repertorio.",
     "Alumna con TDAH que necesita adaptaciones en la gestión del tiempo, descansos pautados y refuerzo de la lectura musical.",
     "Se recomienda fragmentar las tareas, usar señales visuales y coordinar con la tutora del IES para mantener las mismas estrategias."],

    # Carla 10005 — TDAH + ansiedad (informe más antiguo)
    [10005, datetime(2025, 11, 20),
     "Crisis de ansiedad ante evaluaciones formales.",
     "Capacidad técnica adecuada al nivel cuando no hay presión externa.",
     "Tendencia al bloqueo en audiciones y clases magistrales.",
     "Buena interacción con compañeros/as en ensayos de cámara.",
     "Dificultad para generalizar lo aprendido a contextos nuevos.",
     "Creatividad destacable en improvisación y composición libre.",
     "Alumna con TDAH y ansiedad que requiere adaptación de las condiciones de evaluación, apoyo emocional continuado y coordinación con profesional externo.",
     "Se aconseja permitir evaluaciones en entorno controlado, avisar con antelación de las audiciones y mantener comunicación regular con la familia."],

    # Carla 10005 — informe más reciente (debe prevalecer este)
    [10005, datetime(2026, 2, 10),
     "Persisten las crisis de ansiedad, aunque se han reducido en frecuencia.",
     "Mejora notable en la gestión de la frustración y en la relación con el profesorado.",
     "Evita participar voluntariamente en audiciones.",
     "Participa activamente en pequeño grupo; ha mejorado el trabajo cooperativo.",
     "Lectura a primera vista aún por debajo del nivel esperado.",
     "Rendimiento técnico bueno; musicalidad e interpretación expresiva destacadas.",
     "Se mantienen las necesidades del primer informe. Se añade la necesidad de un plan gradual de exposición a situaciones de interpretación pública.",
     "Continuar con el protocolo de evaluación adaptada. Iniciar exposición progresiva a audiciones (primero en grupo, luego individual). Coordinación trimestral con la psicóloga externa."],

    # Sofía 10009 — Dislexia + ansiedad
    [10009, datetime(2025, 12, 5),
     "Material escrito convencional dificulta el acceso a la información musical.",
     "Excelente capacidad auditiva y de imitación.",
     "Se retrae cuando debe leer en voz alta o seguir indicaciones escritas en clase.",
     "Compañerismo y disposición a ayudar a otros.",
     "Lentitud en la lectura de partituras y dificultad con la nomenclatura convencional.",
     "Gran sensibilidad interpretativa y buena respuesta a metodologías prácticas.",
     "Alumna con dislexia y ansiedad asociada que necesita materiales adaptados (tipografía, espaciado, colores) y metodología centrada en lo auditivo-práctico.",
     "Proporcionar partituras con tipografía OpenDyslexic, ampliar los plazos de entrega y priorizar la evaluación práctica sobre la escrita."],

    # Noa 10015 — TEA leve
    [10015, datetime(2026, 1, 18),
     "Dificultad para interpretar instrucciones ambiguas o implícitas.",
     "Constancia, disciplina y alto interés por la música.",
     "Necesita estructura y anticipación para participar en actividades no rutinarias.",
     "Una vez integrada, participa con entusiasmo y aporta ideas originales.",
     "Puede perder el hilo en sesiones con múltiples cambios de actividad.",
     "Excelente capacidad de concentración cuando la tarea es clara y estructurada.",
     "Alumna con TEA leve que requiere anticipación de cambios, instrucciones explícitas y material visual de apoyo.",
     "Usar apoyos visuales (horarios, secuencias), anticipar los cambios de actividad con al menos 5 minutos, y coordinar con la orientadora del IES las adaptaciones."],

    # NIA 10099 — Solo existe en Informes, NO tiene ficha
    [10099, datetime(2026, 2, 20),
     "Dificultad de acceso al transporte adaptado.",
     "Familia implicada y alto interés del alumno/a.",
     "Participación limitada en actividades fuera del aula habitual.",
     "Buena relación con el grupo clase.",
     "Rendimiento irregular por absentismo justificado.",
     "Capacidad de trabajo autónomo cuando dispone de material adaptado.",
     "Alumno/a con necesidades de apoyo por situación sociofamiliar compleja. ¶ Requiere flexibilización horaria y seguimiento de asistencia.",
     "Coordinar con servicios sociales. ¶ Permitir asistencia flexible y facilitar materiales para trabajo autónomo."],

    # Hugo 10002 — Dislexia
    [10002, datetime(2026, 3, 1),
     "Lectura lenta de textos escritos y nomenclatura musical.",
     "Buena comprensión oral y capacidad de razonamiento musical.",
     "Tiende a inhibirse en actividades que implican lectura pública.",
     "Participación activa en debates y actividades orales.",
     "Errores frecuentes en la escritura de dictados musicales.",
     "Memoria auditiva excelente; reproduce fragmentos complejos tras escucharlos.",
     "Alumno con dislexia que necesita adaptaciones en la presentación del material escrito y evaluación preferentemente oral o práctica.",
     "Ampliar tiempos de lectura, usar partituras con mayor espaciado y facilitar grabaciones de audio como complemento al material escrito."],
]


def main():
    wb = Workbook()
    ws_reg = wb.active
    ws_reg.title = "Sheet1"
    ws_reg.append(REG_HEADERS)
    for r in build_registro_rows():
        ws_reg.append(r)

    ws_fic = wb.create_sheet("Ficha")
    ws_fic.append(FICHA_HEADERS)
    for r in ficha_rows:
        ws_fic.append(r)

    ws_inf = wb.create_sheet("Informes")
    ws_inf.append(INFORME_HEADERS)
    for r in informe_rows:
        ws_inf.append(r)

    wb.save(OUT)
    print(f"Generado: {OUT}")
    print(f"  Sheet1: {len(acts)} actuaciones")
    print(f"  Ficha:  {len(ficha_rows)} alumnos")
    print(f"  Informes: {len(informe_rows)} informes")

if __name__ == "__main__":
    main()
